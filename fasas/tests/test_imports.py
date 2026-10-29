import datetime
import io
import zipfile

from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from fasas.audit.models import AuditLog
from fasas.clinical.models import Therapy
from fasas.documents.models import Document
from fasas.guests.models import AdmissionEvent, Guest
from fasas.imports.models import ImportBatch
from fasas.imports.services import run_workbook_import
from fasas.imports.spec import SHEETS
from fasas.imports.template import build_template
from fasas.pharmacy.models import StockLot

from .base import FasasTestCase

NEW_CF = "RSSMRA85T10A562S"


def workbook(rows):
    """Compila il modello ufficiale: rows = {foglio: [{intestazione: valore}]}."""
    wb = load_workbook(io.BytesIO(build_template()))
    for sheet_name, entries in rows.items():
        ws = wb[sheet_name]
        headers = [str(c.value).replace(" *", "") for c in ws[1]]
        for row, entry in enumerate(entries, start=2):
            for col, header in enumerate(headers, start=1):
                ws.cell(row=row, column=col, value=entry.get(header))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sample():
    return workbook({
        "Ospiti": [{"Codice fiscale": NEW_CF, "Cognome": "Bianchi", "Nome": "Anna", "Sesso": "Femmina",
                    "Data di nascita": "10/12/1985", "Data ingresso": datetime.date(2024, 3, 1),
                    "Tag clinici": "Diabetico, Nuovo tag"}],
        "Contatti": [{"Codice fiscale ospite": NEW_CF, "Tipo": "Amministratore di sostegno", "Nome e cognome": "Luca Neri",
                      "Telefono": "333 1234567", "Da contattare in emergenza": "SI"}],
        "Diagnosi": [{"Codice fiscale ospite": NEW_CF, "Descrizione": "Diabete mellito tipo 2", "Principale": "sì"}],
        "Allergie": [{"Codice fiscale ospite": NEW_CF, "Sostanza / principio attivo": "amoxicillina",
                      "Gravità": "Grave / anafilassi"}],
        "Terapie": [
            {"Codice fiscale ospite": NEW_CF, "Farmaco (nome commerciale)": "Metformina EG", "Principio attivo": "metformina",
             "Dosaggio confezione": "500 mg", "Dose": "1", "Orari": "8, 20.00", "Inizio": "01/03/2024",
             "Medico prescrittore": "Dott. Verdi"},
            {"Codice fiscale ospite": NEW_CF, "Farmaco (nome commerciale)": "Tachipirina", "Principio attivo": "paracetamolo",
             "Dosaggio confezione": "500 mg", "Dose": "1", "Al bisogno (SOS)": "SI", "Indicazione SOS": "TC > 38",
             "Inizio": "01/03/2024", "Medico prescrittore": "Dott. Verdi"},
        ],
        "Magazzino": [{"Farmaco (nome commerciale)": "Metformina EG", "Dosaggio confezione": "500 mg", "Lotto": "A1",
                       "Scadenza": "31/12/2030", "Quantità": "60"}],
    })


class WorkbookImportTests(FasasTestCase):
    def test_template_has_every_sheet_and_required_headers(self):
        wb = load_workbook(io.BytesIO(build_template()))
        for sheet in SHEETS:
            headers = [c.value for c in wb[sheet.name][1]]
            self.assertIn("Codice fiscale *" if sheet.name == "Ospiti" else headers[0], headers)
            self.assertEqual(len(headers), len(sheet.columns))

    def test_preview_saves_nothing(self):
        report = run_workbook_import(sample(), self.coordinator, commit=False)
        self.assertEqual(report.errors, [])
        self.assertFalse(report.committed)
        self.assertEqual(report.counts["Ospiti"]["created"], 1)
        self.assertEqual(report.counts["Terapie"]["created"], 2)
        self.assertFalse(Guest.objects.filter(fiscal_code=NEW_CF).exists())

    def test_import_creates_records_with_audit_and_can_be_repeated(self):
        report = run_workbook_import(sample(), self.coordinator, commit=True, filename="fasas.xlsx")
        self.assertTrue(report.committed, report.issues)
        guest = Guest.objects.get(fiscal_code=NEW_CF)
        self.assertEqual(guest.admission_date, datetime.date(2024, 3, 1))
        self.assertEqual(guest.birth_date, datetime.date(1985, 12, 10))
        self.assertEqual(sorted(t.name for t in guest.tags.all()), ["Diabetico", "Nuovo tag"])
        self.assertTrue(guest.contacts.get().is_emergency_contact)
        self.assertTrue(guest.diagnoses.get().is_primary)
        self.assertEqual(guest.allergies.get().severity, "severe")
        therapy = Therapy.objects.get(guest=guest, is_sos=False)
        self.assertEqual(therapy.schedule_times, ["08:00", "20:00"])
        self.assertTrue(Therapy.objects.get(guest=guest, is_sos=True).sos_indication)
        self.assertEqual(StockLot.objects.get(lot_number="A1").quantity, 60)
        self.assertTrue(AdmissionEvent.objects.filter(guest=guest, kind="admission").exists())
        self.assertTrue(AuditLog.objects.filter(guest_id=guest.pk, source="import").exists())

        again = run_workbook_import(sample(), self.coordinator, commit=True)
        self.assertEqual(again.total_created, 0)
        self.assertEqual(again.counts["Terapie"]["skipped"], 2)
        self.assertEqual(Therapy.objects.filter(guest=guest).count(), 2)

    def test_any_error_blocks_the_whole_import(self):
        data = workbook({
            "Ospiti": [{"Codice fiscale": NEW_CF, "Cognome": "Bianchi", "Nome": "Anna", "Data di nascita": "31/02/1985",
                        "Data ingresso": "01/03/2024"},
                       {"Codice fiscale": "VRDLCU70A01H501U", "Cognome": "Verdi", "Nome": "Luca",
                        "Data di nascita": "01/01/1970", "Data ingresso": "01/03/2024"}],
            "Terapie": [{"Codice fiscale ospite": "XXXXXX00X00X000X", "Farmaco (nome commerciale)": "Depakin",
                         "Principio attivo": "valproato", "Dose": "1", "Orari": "08:00", "Inizio": "01/03/2024",
                         "Medico prescrittore": "Dott. Verdi"}],
        })
        report = run_workbook_import(data, self.coordinator, commit=True)
        self.assertFalse(report.committed)
        messages = " ".join(i.message for i in report.errors)
        self.assertIn("data non valida", messages)
        self.assertIn("non trovato", messages)
        self.assertFalse(Guest.objects.filter(fiscal_code="VRDLCU70A01H501U").exists())

    def test_conflicts_with_allergies_are_reported(self):
        self.guest.allergies.create(substance="paracetamolo", severity="severe")
        data = workbook({"Terapie": [{"Codice fiscale ospite": self.guest.fiscal_code,
                                      "Farmaco (nome commerciale)": "Tachipirina", "Principio attivo": "paracetamolo",
                                      "Dose": "1", "Orari": "08:00", "Inizio": "01/03/2024",
                                      "Medico prescrittore": "Dott. Verdi"}]})
        report = run_workbook_import(data, self.coordinator, commit=True)
        self.assertTrue(report.committed)
        self.assertTrue(any("ALLERGIA" in w.message for w in report.warnings))
        self.assertIn("ALLERGIA", Therapy.objects.get(guest=self.guest).conflict_acknowledged)

    def test_web_flow_preview_then_confirm(self):
        self.login(self.coordinator)
        upload = SimpleUploadedFile("fasas.xlsx", sample())
        response = self.client.post(reverse("imports:workbook"), {"file": upload})
        batch = ImportBatch.objects.get()
        self.assertRedirects(response, reverse("imports:batch", args=[batch.pk]))
        page = self.client.get(reverse("imports:batch", args=[batch.pk]))
        self.assertContains(page, "Conferma importazione")
        stored = batch.payload[0]
        self.assertTrue(default_storage.exists(stored))
        self.client.post(reverse("imports:confirm", args=[batch.pk]))
        batch.refresh_from_db()
        self.assertEqual(batch.status, "completed")
        self.assertEqual(batch.payload, [])
        self.assertFalse(default_storage.exists(stored))
        self.assertTrue(Guest.objects.filter(fiscal_code=NEW_CF).exists())

    def test_import_is_for_coordinators_only(self):
        for user in (self.operator, self.doctor, self.admin):
            self.login(user)
            self.assertEqual(self.client.get(reverse("imports:home")).status_code, 403, user.username)
        self.login(self.coordinator)
        self.assertEqual(self.client.get(reverse("imports:home")).status_code, 200)
        response = self.client.get(reverse("imports:template"))
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class DocumentImportTests(FasasTestCase):
    def test_zip_files_are_matched_by_fiscal_code_or_name(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr(f"{self.guest.fiscal_code} consenso informato.pdf", b"%PDF-1.4 uno")
            archive.writestr("Rossi Mario/diario_2024.pdf", b"%PDF-1.4 due")
            archive.writestr("sconosciuto.pdf", b"%PDF-1.4 tre")
            archive.writestr("__MACOSX/._x.pdf", b"x")
            archive.writestr("programma.exe", b"MZ")
        self.login(self.coordinator)
        upload = SimpleUploadedFile("scansioni.zip", buffer.getvalue(), content_type="application/zip")
        self.client.post(reverse("imports:documents"), {"files": [upload], "category": "medical_report"})
        batch = ImportBatch.objects.get(kind="documents")
        by_path = {e["path"]: e for e in batch.payload}
        self.assertEqual(len(by_path), 4)
        self.assertEqual(by_path[f"{self.guest.fiscal_code} consenso informato.pdf"]["guest_id"], self.guest.pk)
        self.assertEqual(by_path[f"{self.guest.fiscal_code} consenso informato.pdf"]["title"], "Consenso informato")
        self.assertEqual(by_path["Rossi Mario/diario_2024.pdf"]["guest_id"], self.guest.pk)
        self.assertIsNone(by_path["sconosciuto.pdf"]["guest_id"])
        self.assertIn("non ammessa", by_path["programma.exe"]["problem"])

        post = {}
        for index, entry in enumerate(batch.payload):
            if not entry["problem"]:
                post[f"include_{index}"] = "on"
                post[f"guest_{index}"] = str(entry["guest_id"] or "")
                post[f"title_{index}"] = entry["title"]
        self.client.post(reverse("imports:confirm", args=[batch.pk]), post)
        batch.refresh_from_db()
        self.assertEqual(batch.summary["created"], 2)
        docs = Document.objects.filter(guest=self.guest, category="medical_report")
        self.assertEqual(docs.count(), 2)
        self.assertTrue(all(not e["storage"] for e in batch.payload))
