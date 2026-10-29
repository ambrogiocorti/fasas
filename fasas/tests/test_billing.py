import datetime
import io
from decimal import Decimal

from django.urls import reverse
from openpyxl import load_workbook

from fasas.accounts.models import Role
from fasas.audit.context import audit_context
from fasas.billing import services
from fasas.billing.exports import attendance_workbook, statement_pdf
from fasas.billing.models import AttendanceMonth, BillingSettings, FeeStatement, GuestFeeLine
from fasas.guests.models import AdmissionEvent

from .base import FasasTestCase


class BillingTests(FasasTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.clerk = cls.make_user("amm", Role.ADMINISTRATIVE)

    def setUp(self):
        self.first = services.shift_month(services.month_start(), -2)  # un mese sicuramente concluso
        self.guest.admission_date = self.first - datetime.timedelta(days=60)
        self.guest.save()
        AdmissionEvent.objects.create(guest=self.guest, kind="admission", date=self.guest.admission_date)
        self.day = lambda n: self.first.replace(day=n)

    def move(self, kind, day):
        AdmissionEvent.objects.create(guest=self.guest, kind=kind, date=self.day(day))

    def days_of(self, guest=None):
        rows = {g.pk: values for g, values, _, _ in services.compute_month(self.first)}
        return rows[(guest or self.guest).pk]

    def fee(self, **extra):
        data = {"guest": self.guest, "description": "Quota sociale", "payer": "municipality",
                "payer_name": "Comune di Prova", "daily_rate": Decimal("42.00"), "valid_from": self.guest.admission_date}
        data.update(extra)
        with audit_context(user=self.clerk):
            return GuestFeeLine.objects.create(**data)

    def test_days_follow_movements_and_rules(self):
        self.move("hospitalization", 8)
        self.move("return", 14)
        self.move("absence", 20)
        self.move("return", 23)
        days = self.days_of()
        self.assertEqual(days[7], "P")
        self.assertEqual(days[8], "P")  # giorno di uscita conteggiato come presenza
        self.assertEqual([days[d] for d in range(9, 14)], ["H"] * 5)
        self.assertEqual(days[14], "P")  # giorno di rientro
        self.assertEqual([days[d] for d in (21, 22)], ["F", "F"])
        rules = BillingSettings.get()
        rules.departure_day_present = False
        rules.return_day_present = False
        rules.save()
        days = self.days_of()
        self.assertEqual(days[8], "H")
        self.assertEqual(days[14], "H")
        self.assertEqual(days[23], "F")

    def test_admission_and_discharge_days(self):
        other = self.make_guest("VRDLCU70A01H501U", first_name="Luca", last_name="Verdi", admission_date=self.day(10))
        AdmissionEvent.objects.create(guest=other, kind="admission", date=self.day(10))
        AdmissionEvent.objects.create(guest=other, kind="discharge", date=self.day(20))
        other.discharge_date = self.day(20)
        other.save()
        days = self.days_of(other)
        self.assertEqual(days[9], "-")
        self.assertEqual(days[10], "P")
        self.assertEqual(days[19], "P")
        self.assertEqual(days[20], "-")  # giorno di dimissione non conteggiato per default
        self.assertEqual(days[25], "-")

    def test_correction_close_and_reopen(self):
        services.save_correction(self.guest, self.day(3), "A", "Uscita non registrata", self.clerk)
        self.assertEqual(self.days_of()[3], "A")
        with self.assertRaises(services.BillingError):
            services.close_month(services.month_start(), self.clerk)  # mese in corso
        month = services.close_month(self.first, self.clerk, "Tutto regolare")
        self.assertTrue(month.is_closed)
        row = month.rows.get(guest=self.guest)
        self.assertEqual(row.absent, 1)
        # Dopo la chiusura i movimenti non cambiano più le presenze del mese
        self.move("hospitalization", 10)
        self.assertEqual(services.attendance(self.first)[1][0][1][12], "P")
        with self.assertRaises(services.BillingError):
            services.save_correction(self.guest, self.day(4), "A", "x", self.clerk)
        services.reopen_month(month, self.clerk, "Ricovero registrato in ritardo")
        self.assertEqual(services.attendance(self.first)[1][0][1][12], "H")

    def test_statements_amounts_numbering_and_payments(self):
        self.move("hospitalization", 8)
        self.move("return", 14)  # 9-13: 5 giorni di ricovero
        self.fee(hospital_percent=70)
        self.fee(description="Compartecipazione", payer="guest", payer_name="Elena Rossi", daily_rate=Decimal("15"),
                 hospital_percent=100, max_absence_days=3)
        month = services.close_month(self.first, self.clerk)
        created, skipped, missing = services.generate_statements(month, self.clerk)
        self.assertEqual((len(created), skipped, missing), (2, 0, []))
        total_days = len(services.month_days(self.first))
        social = FeeStatement.objects.get(payer_name="Comune di Prova")
        expected = Decimal(total_days - 5) * Decimal("42") + Decimal("5") * Decimal("42") * Decimal("0.7")
        self.assertEqual(social.total, expected.quantize(Decimal("0.01")))
        family = FeeStatement.objects.get(payer_name="Elena Rossi")
        capped = [line for line in family.lines if line["status"] == "X"]
        self.assertEqual(capped[0]["days"], 2)  # assenze oltre il limite di 3 giorni
        self.assertEqual(family.total, Decimal(total_days - 2) * 15)

        first = services.issue_statement(social, self.clerk)
        second = services.issue_statement(family, self.clerk)
        year = first.issue_date.year
        self.assertEqual((first.number, second.number), (f"{year}/0001", f"{year}/0002"))
        self.assertEqual(first.due_date, first.issue_date + datetime.timedelta(days=30))
        with self.assertRaises(services.BillingError):
            services.reopen_month(month, self.clerk, "errore")
        services.register_payment(first, first.total - 10, first.issue_date, "transfer", "CRO 1", self.clerk)
        first.refresh_from_db()
        self.assertEqual(first.status, "issued")
        services.register_payment(first, Decimal("10"), first.issue_date, "transfer", "CRO 2", self.clerk)
        first.refresh_from_db()
        self.assertEqual(first.status, "paid")
        with self.assertRaises(services.BillingError):
            services.cancel_statement(first, self.clerk, "errore")
        services.cancel_statement(second, self.clerk, "Intestatario errato")
        second.refresh_from_db()
        self.assertEqual(second.status, "cancelled")
        # Rigenerando, i prospetti già emessi non vengono duplicati; quello annullato si può rifare
        created, skipped, _ = services.generate_statements(month, self.clerk)
        self.assertEqual((len(created), skipped), (1, 1))

    def test_exports(self):
        self.fee()
        month = services.close_month(self.first, self.clerk)
        statement = services.issue_statement(services.generate_statements(month, self.clerk)[0][0], self.clerk)
        wb = load_workbook(io.BytesIO(attendance_workbook(self.first)))
        sheet = wb.worksheets[0]
        self.assertEqual(sheet.cell(row=3, column=2).value, self.guest.fiscal_code)
        self.assertTrue(statement_pdf(statement, self.clerk).startswith(b"%PDF"))

    def test_role_access(self):
        self.login(self.clerk)
        response = self.client.get(reverse("core:dashboard"))
        self.assertContains(response, "<h1>Amministrazione")
        for name in ("billing:attendance", "billing:fees", "billing:statements", "billing:settings"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        for name, args in (("guests:list", []), ("guests:detail", [self.guest.pk]), ("clinical:board", []),
                           ("accounts:user_list", [])):
            self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 403, name)
        self.login(self.coordinator)
        self.assertEqual(self.client.get(reverse("billing:attendance")).status_code, 200)
        for name in ("billing:fees", "billing:statements"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)
        for user in (self.operator, self.doctor, self.admin):
            self.login(user)
            self.assertEqual(self.client.get(reverse("billing:attendance")).status_code, 403, user.username)

    def test_web_flow_close_generate_issue(self):
        self.fee()
        self.login(self.clerk)
        mese = f"{self.first:%Y-%m}"
        page = self.client.get(reverse("billing:attendance"), {"mese": mese})
        self.assertContains(page, "Chiudi il mese")
        self.client.post(reverse("billing:close_month"), {"mese": mese, "confirm": "on"})
        self.assertTrue(AttendanceMonth.objects.get(month=self.first).is_closed)
        self.client.post(reverse("billing:generate"), {"mese": mese})
        self.client.post(reverse("billing:issue_all"), {"mese": mese})
        statement = FeeStatement.objects.get()
        self.assertEqual(statement.status, "issued")
        response = self.client.post(statement.get_absolute_url(), {
            "date": statement.issue_date.isoformat(), "amount": str(statement.total), "method": "transfer",
            "reference": "CRO"})
        self.assertRedirects(response, statement.get_absolute_url())
        statement.refresh_from_db()
        self.assertEqual(statement.status, "paid")
