import datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from fasas.backups.crypto import BackupCryptoError, decrypt_stream, encrypt_stream
from fasas.backups.services import create_backup, verify_backup
from fasas.clinical.models import ClinicalEvent
from fasas.core.encryption import read_encrypted
from fasas.reports.guest_export import build_guest_folder
from fasas.reports.models import GeneratedReport, ReportSchedule
from fasas.reports.services import generate_report, purge_expired_reports, run_due_schedules

from .base import FasasTestCase


class ReportTests(FasasTestCase):
    def setUp(self):
        self.event = ClinicalEvent.objects.create(guest=self.guest, event_type="fall", description="=HYPERLINK(\"x\")")

    def test_pdf_and_excel_reports(self):
        # i report vanno per giornata di turno: dopo mezzanotte l'evento è della notte del giorno prima
        today = self.event.shift_date
        pdf = generate_report("events", "pdf", today, today, user=self.coordinator)
        data = read_encrypted(pdf.storage_name)
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertIn("Eventi", PdfReader(BytesIO(data)).pages[0].extract_text())
        xlsx = generate_report("events", "xlsx", today, today, user=self.coordinator)
        ws = load_workbook(BytesIO(read_encrypted(xlsx.storage_name))).active
        values = [c.value for row in ws.iter_rows() for c in row if c.value]
        self.assertTrue(any(str(v).startswith("'=HYPERLINK") for v in values))  # niente formule iniettate
        self.assertEqual(pdf.expires_on.year, timezone.localdate().year + 5)

    def test_event_after_midnight_belongs_to_previous_night(self):
        tz = timezone.get_current_timezone()
        night = timezone.make_aware(datetime.datetime(2026, 11, 1, 0, 15), tz)
        event = ClinicalEvent.objects.create(guest=self.guest, event_type="fall", description="caduta notturna",
                                             occurred_at=night)
        self.assertEqual(event.shift_date, datetime.date(2026, 10, 31))
        report = generate_report("events", "xlsx", event.shift_date, event.shift_date, user=self.coordinator)
        ws = load_workbook(BytesIO(read_encrypted(report.storage_name))).active
        self.assertIn("caduta notturna", [c.value for row in ws.iter_rows() for c in row])

    def test_all_report_types_render(self):
        today = timezone.localdate()
        for report_type in ["daily_register", "medications", "events", "parameters", "protocols"]:
            for fmt in ["pdf", "xlsx"]:
                generate_report(report_type, fmt, today - datetime.timedelta(days=7), today, guest=self.guest)
        self.assertEqual(GeneratedReport.objects.count(), 10)

    def test_scheduled_reports_once_per_period_and_retention(self):
        ReportSchedule.objects.create(name="Eventi mensili", report_type="events", frequency="monthly", retention_years=1)
        self.assertEqual(run_due_schedules()["generated"], 1)
        self.assertEqual(run_due_schedules()["generated"], 0)
        report = GeneratedReport.objects.get()
        self.assertEqual(report.expires_on.year, timezone.localdate().year + 1)
        self.assertEqual(purge_expired_reports(today=report.expires_on + datetime.timedelta(days=1))["removed"], 1)
        self.assertFalse(GeneratedReport.objects.exists())

    def test_guest_folder_export(self):
        from fasas.documents.services import create_document

        from .test_documents import sample_pdf

        create_document(self.guest, "medical_report", "Relazione", user=self.coordinator, data=sample_pdf(),
                        filename="r.pdf")
        pdf = build_guest_folder(self.guest, self.coordinator, include_documents=True)
        reader = PdfReader(BytesIO(pdf))
        text = "".join(page.extract_text() for page in reader.pages)
        self.assertIn("Dati anagrafici", text)
        self.assertIn("Relazione di prova", text)  # allegato accodato
        self.login(self.coordinator)
        response = self.client.post(f"/report/ospite/{self.guest.pk}/cartella/", {"include_documents": "on"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(self.client.get(response["Location"]).content.startswith(b"%PDF"))

    def test_operator_cannot_export(self):
        self.login(self.operator)
        self.assertEqual(self.client.get(f"/report/ospite/{self.guest.pk}/cartella/").status_code, 403)


class BackupTests(FasasTestCase):
    def test_stream_encryption_detects_truncation(self):
        payload = b"x" * (4 * 1024 * 1024 + 123)
        encrypted = BytesIO()
        encrypt_stream(BytesIO(payload), encrypted)
        out = BytesIO()
        decrypt_stream(BytesIO(encrypted.getvalue()), out)
        self.assertEqual(out.getvalue(), payload)
        first_chunk_only = encrypted.getvalue()[: 16 + 4 + 4 * 1024 * 1024 + 16]
        with self.assertRaises(BackupCryptoError):
            decrypt_stream(BytesIO(first_chunk_only), BytesIO())
        tampered = bytearray(encrypted.getvalue())
        tampered[100] ^= 1
        with self.assertRaises(BackupCryptoError):
            decrypt_stream(BytesIO(bytes(tampered)), BytesIO())

    def test_backup_create_and_verify(self):
        record = create_backup(user=self.admin)
        self.assertEqual(record.status, "ok", record.error)
        self.assertTrue(verify_backup(record))
        with open(record.path, "r+b") as handle:
            handle.seek(200)
            byte = handle.read(1)
            handle.seek(200)
            handle.write(bytes([byte[0] ^ 1]))
        self.assertFalse(verify_backup(record))
