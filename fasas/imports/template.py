"""Modello Excel da compilare con i dati del FASAS cartaceo."""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .spec import SHEETS, parse_date

ROWS = 500
HEADER_FILL = PatternFill("solid", fgColor="1F5FAD")
REQUIRED_FILL = PatternFill("solid", fgColor="B3261E")
LIST_SHEET = "Elenchi"

INSTRUCTIONS = [
    "Come compilare il modello",
    "",
    "1. Scegliete una data di passaggio (es. il primo del mese): da quel giorno si registra solo su FASAS Digitale.",
    "2. Compilate un foglio alla volta, una riga per voce. Le colonne in rosso sono obbligatorie.",
    "3. Il codice fiscale collega i fogli: ogni contatto, diagnosi, allergia e terapia va riferita al codice "
    "fiscale dell'ospite.",
    "4. Nelle colonne con l'elenco a scelta usate solo i valori proposti. Date nel formato GG/MM/AAAA.",
    "5. Terapie: riportate solo quelle in corso, dal foglio unico di terapia aggiornato e firmato dal medico.",
    "6. Lo storico (diari, parametri, eventi degli anni passati) non va ricopiato: si scansiona e si carica "
    "come documento (Importa dati → Documenti scansionati).",
    "7. Caricate il file: il sistema mostra un'anteprima con gli errori riga per riga senza salvare nulla. "
    "Correggete e ricaricate finché non ci sono errori, poi confermate.",
    "8. Si può ricaricare lo stesso file più volte: le voci già presenti vengono riconosciute e saltate.",
    "",
    "Dopo l'importazione un medico deve verificare le terapie importate confrontandole con il cartaceo.",
    "",
    "Contenuto dei fogli",
]


def build_template():
    wb = Workbook()
    info = wb.active
    info.title = "Istruzioni"
    info.column_dimensions["A"].width = 120
    for line in INSTRUCTIONS:
        info.append([line])
    for sheet in SHEETS:
        info.append([f"• {sheet.name}: {sheet.description}"])
    info["A1"].font = Font(bold=True, size=14)
    info[f"A{INSTRUCTIONS.index('Contenuto dei fogli') + 1}"].font = Font(bold=True)

    lists = wb.create_sheet(LIST_SHEET)
    list_col = 0
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.name)
        ws.freeze_panes = "A2"
        for index, column in enumerate(sheet.columns, start=1):
            letter = get_column_letter(index)
            cell = ws.cell(row=1, column=index, value=column.header + (" *" if column.required else ""))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = REQUIRED_FILL if column.required else HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if column.help:
                cell.comment = Comment(column.help, "FASAS")
            ws.column_dimensions[letter].width = column.width
            if column.parser is parse_date:
                for row in range(2, ROWS + 2):
                    ws.cell(row=row, column=index).number_format = "DD/MM/YYYY"
            if column.labels:
                list_col += 1
                list_letter = get_column_letter(list_col)
                lists.cell(row=1, column=list_col, value=f"{sheet.name} – {column.header}").font = Font(bold=True)
                for row, label in enumerate(column.labels, start=2):
                    lists.cell(row=row, column=list_col, value=label)
                source = f"{LIST_SHEET}!${list_letter}$2:${list_letter}${len(column.labels) + 1}"
                validation = DataValidation(type="list", formula1=f"={source}", allow_blank=True,
                                            showErrorMessage=True, errorTitle="Valore non previsto",
                                            error="Scegliere un valore dall'elenco.")
                validation.add(f"{letter}2:{letter}{ROWS + 1}")
                ws.add_data_validation(validation)
        ws.row_dimensions[1].height = 42
    lists.sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
