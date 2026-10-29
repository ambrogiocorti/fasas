"""
Importazione del FASAS cartaceo da un file Excel compilato sul modello.

L'importazione è «tutto o niente»: si esegue sempre dentro una transazione; in anteprima
(o se c'è anche un solo errore) la transazione viene annullata, quindi l'anteprima applica
esattamente gli stessi controlli dell'importazione vera. Ogni dato creato finisce nell'audit
log con origine «import». Le voci già presenti vengono riconosciute e saltate, così lo stesso
file può essere ricaricato dopo aver corretto gli errori.
"""
import io
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.files.storage import default_storage
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook

from fasas.audit.context import audit_context
from fasas.clinical.models import Therapy
from fasas.clinical.services import check_therapy_conflicts
from fasas.core.encryption import read_encrypted, sha256, store_encrypted
from fasas.core.models import Community
from fasas.guests.models import AdmissionEvent, Allergy, ClinicalTag, Diagnosis, Guest, GuestContact
from fasas.pharmacy.models import Drug, StockLot
from fasas.pharmacy.services import load_lot

from .models import ImportBatch
from .spec import SHEET_BY_NAME, SHEETS, CellError

SHEET_PERMISSIONS = {
    "Ospiti": "guests.add_guest",
    "Contatti": "guests.add_guestcontact",
    "Diagnosi": "guests.add_diagnosis",
    "Allergie": "guests.add_allergy",
    "Terapie": "clinical.add_therapy",
    "Magazzino": "pharmacy.add_stocklot",
}

# Carattere di controllo del codice fiscale (DM 23/12/1976)
_ODD = dict(zip("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ",
                [1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 2, 4, 18, 20, 11, 3, 6, 8,
                 12, 14, 16, 10, 22, 25, 24, 23], strict=True))
_EVEN = {c: (int(c) if c.isdigit() else ord(c) - ord("A")) for c in "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"}


def fiscal_code_checksum_ok(code):
    code = (code or "").upper()
    if len(code) != 16 or not code.isalnum():
        return False
    total = sum(_ODD[c] if i % 2 == 0 else _EVEN[c] for i, c in enumerate(code[:15]))
    return chr(ord("A") + total % 26) == code[15]


@dataclass
class Issue:
    sheet: str
    row: int
    column: str
    message: str
    level: str = "error"

    def as_dict(self):
        return {"sheet": self.sheet, "row": self.row, "column": self.column, "message": self.message,
                "level": self.level}


@dataclass
class ImportReport:
    issues: list = field(default_factory=list)
    counts: dict = field(default_factory=lambda: {s.name: {"created": 0, "skipped": 0} for s in SHEETS})
    committed: bool = False

    def error(self, sheet, row, message, column=""):
        self.issues.append(Issue(sheet, row, column, message))

    def warning(self, sheet, row, message, column=""):
        self.issues.append(Issue(sheet, row, column, message, "warning"))

    def created(self, sheet):
        self.counts[sheet]["created"] += 1

    def skipped(self, sheet):
        self.counts[sheet]["skipped"] += 1

    @property
    def errors(self):
        return [i for i in self.issues if i.level == "error"]

    @property
    def warnings(self):
        return [i for i in self.issues if i.level == "warning"]

    @property
    def total_created(self):
        return sum(c["created"] for c in self.counts.values())

    def as_dict(self):
        return {"counts": self.counts, "issues": [i.as_dict() for i in self.issues], "committed": self.committed}


class _Rollback(Exception):
    pass


def _norm_header(text):
    return " ".join(str(text or "").replace("*", " ").lower().split())


def read_workbook(data, report):
    """Restituisce {nome foglio: [(numero riga, valori)]} con i valori già convertiti."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        report.error("File", 0, "Il file non è un foglio Excel (.xlsx) leggibile.")
        return {}
    parsed = {}
    for ws in wb.worksheets:
        sheet = SHEET_BY_NAME.get(ws.title.strip().lower())
        if sheet is None:
            continue
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or ()
        positions = {}
        for index, value in enumerate(header):
            for column in sheet.columns:
                if _norm_header(value) == _norm_header(column.header):
                    positions[column.key] = index
        missing = [c.header for c in sheet.columns if c.required and c.key not in positions]
        if missing:
            report.error(sheet.name, 1, f"Mancano le colonne obbligatorie: {', '.join(missing)}. "
                                        "Usare il modello scaricato dal sistema.")
            continue
        entries = []
        for number, row in enumerate(rows, start=2):
            if not row or all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            values, ok = {}, True
            for column in sheet.columns:
                raw = row[positions[column.key]] if column.key in positions and positions[column.key] < len(row) else None
                try:
                    value = column.parse(raw)
                except CellError as exc:
                    report.error(sheet.name, number, str(exc), column.header)
                    ok = False
                    continue
                if column.required and value in (None, "", []):
                    report.error(sheet.name, number, "campo obbligatorio", column.header)
                    ok = False
                values[column.key] = value
            if ok:
                entries.append((number, values))
        parsed[sheet.name] = entries
    wb.close()
    if not parsed and not report.errors:
        report.error("File", 0, "Nessun foglio riconosciuto: usare il modello scaricato dal sistema.")
    return parsed


def _validation_messages(exc):
    if hasattr(exc, "message_dict"):
        return [(name if name != "__all__" else "", " ".join(msgs)) for name, msgs in exc.message_dict.items()]
    return [("", " ".join(exc.messages))]


def _clean_or_report(instance, report, sheet, row, exclude=None, labels=None):
    try:
        instance.full_clean(exclude=exclude)
    except ValidationError as exc:
        for name, message in _validation_messages(exc):
            report.error(sheet, row, message, (labels or {}).get(name, name))
        return False
    return True


class _Importer:
    def __init__(self, parsed, user, report):
        self.parsed, self.user, self.report = parsed, user, report
        cfs = {v.get("fiscal_code") or v.get("guest_cf") or v.get("owner_cf") for rows in parsed.values()
               for _, v in rows}
        self.guests = {g.fiscal_code: g for g in Guest.objects.filter(fiscal_code__in=[c for c in cfs if c])}
        self.new_drugs = {}

    def run(self):
        for sheet in SHEETS:
            rows = self.parsed.get(sheet.name, [])
            if not rows:
                continue
            perm = SHEET_PERMISSIONS[sheet.name]
            if not self.user.has_perm(perm):
                self.report.error(sheet.name, 0, "Non hai il permesso di importare questo foglio.")
                continue
            handler = getattr(self, f"import_{sheet.name.lower()}")
            for number, values in rows:
                try:
                    with transaction.atomic():
                        handler(sheet, number, values)
                except (IntegrityError, ValidationError, ValueError) as exc:
                    self.report.error(sheet.name, number, f"riga non importabile: {exc}")

    # --- helper ------------------------------------------------------------------------

    def guest_for(self, sheet, row, cf, column="Codice fiscale ospite"):
        guest = self.guests.get(cf)
        if guest is None:
            self.report.error(sheet.name, row, f"ospite con codice fiscale {cf} non trovato "
                                               "(né nel foglio Ospiti né nel sistema)", column)
        elif guest.is_archived:
            self.report.error(sheet.name, row, f"la cartella di {guest} è archiviata", column)
            return None
        return guest

    def drug_for(self, sheet, row, values):
        key = (values["drug_name"].lower(), (values.get("strength") or "").lower())
        if key in self.new_drugs:
            return self.new_drugs[key]
        drug = Drug.objects.filter(name__iexact=values["drug_name"], strength__iexact=values.get("strength") or "").first()
        if drug is None:
            if not values.get("active_ingredient"):
                self.report.error(sheet.name, row, "farmaco non in anagrafica: indicare il principio attivo",
                                  "Principio attivo")
                return None
            drug = Drug(name=values["drug_name"], active_ingredient=values["active_ingredient"],
                        strength=values.get("strength") or "", form=values.get("form") or Drug.Form.TABLET,
                        stock_unit=values.get("stock_unit") or "unità")
            if not _clean_or_report(drug, self.report, sheet.name, row):
                return None
            drug.save()
            self.report.warning(sheet.name, row, f"farmaco «{drug}» aggiunto all'anagrafica")
        elif values.get("active_ingredient") and values["active_ingredient"].lower() != drug.active_ingredient.lower():
            self.report.warning(sheet.name, row, f"in anagrafica «{drug}» ha principio attivo "
                                                 f"«{drug.active_ingredient}»: viene usato quello", "Principio attivo")
        self.new_drugs[key] = drug
        return drug

    # --- fogli -------------------------------------------------------------------------

    def import_ospiti(self, sheet, row, v):
        cf = v["fiscal_code"]
        if cf in self.guests:
            if getattr(self.guests[cf], "_imported_row", None):
                self.report.error(sheet.name, row, f"codice fiscale ripetuto (già alla riga {self.guests[cf]._imported_row})",
                                  "Codice fiscale")
            else:
                self.report.skipped(sheet.name)
                self.report.warning(sheet.name, row, f"{self.guests[cf]} è già presente: anagrafica non modificata")
            return
        if not fiscal_code_checksum_ok(cf) and not cf.startswith(("STP", "ENI")):
            self.report.warning(sheet.name, row, "il carattere di controllo del codice fiscale non torna: "
                                                 "verificare che sia scritto correttamente", "Codice fiscale")
        if v["birth_date"] and v["birth_date"] > timezone.localdate():
            self.report.error(sheet.name, row, "data di nascita nel futuro", "Data di nascita")
            return
        if v["admission_date"] and v["birth_date"] and v["admission_date"] < v["birth_date"]:
            self.report.error(sheet.name, row, "data di ingresso precedente alla nascita", "Data ingresso")
            return
        fields = {k: val for k, val in v.items() if k != "tags"}
        guest = Guest(community=Community.default(), **fields)
        labels = {c.key: c.header for c in sheet.columns}
        if not _clean_or_report(guest, self.report, sheet.name, row, labels=labels):
            return
        guest.save()
        AdmissionEvent.objects.create(guest=guest, kind=AdmissionEvent.Kind.ADMISSION, date=guest.admission_date,
                                      notes="Ingresso registrato dall'importazione del FASAS cartaceo")
        for name in v["tags"]:
            tag = ClinicalTag.objects.filter(name__iexact=name).first()
            if tag is None:
                tag = ClinicalTag.objects.create(name=name[:60])
                self.report.warning(sheet.name, row, f"nuovo tag clinico «{tag}» creato: completarne colore e avvisi",
                                    "Tag clinici")
            guest.tags.add(tag)
        guest._imported_row = row
        self.guests[cf] = guest
        self.report.created(sheet.name)

    def _simple(self, sheet, row, v, model, exists_filter):
        guest = self.guest_for(sheet, row, v["guest_cf"])
        if guest is None:
            return
        fields = {k: val for k, val in v.items() if k != "guest_cf"}
        if model.objects.filter(guest=guest, **exists_filter(fields)).exists():
            self.report.skipped(sheet.name)
            return
        obj = model(guest=guest, **fields)
        labels = {c.key: c.header for c in sheet.columns}
        if _clean_or_report(obj, self.report, sheet.name, row, labels=labels):
            obj.save()
            self.report.created(sheet.name)

    def import_contatti(self, sheet, row, v):
        self._simple(sheet, row, v, GuestContact, lambda f: {"name__iexact": f["name"], "kind": f["kind"]})

    def import_diagnosi(self, sheet, row, v):
        self._simple(sheet, row, v, Diagnosis, lambda f: {"description__iexact": f["description"]})

    def import_allergie(self, sheet, row, v):
        self._simple(sheet, row, v, Allergy, lambda f: {"substance__iexact": f["substance"],
                                                        "category": f["category"]})

    def import_terapie(self, sheet, row, v):
        guest = self.guest_for(sheet, row, v["guest_cf"])
        if guest is None:
            return
        drug = self.drug_for(sheet, row, v)
        if drug is None:
            return
        if v["is_sos"] and v["schedule_times"]:
            self.report.warning(sheet.name, row, "farmaco al bisogno: gli orari indicati non vengono usati", "Orari")
            v["schedule_times"] = []
        same = Therapy.objects.filter(guest=guest, drug=drug, status=Therapy.Status.ACTIVE, is_sos=v["is_sos"],
                                      dose=v["dose"])
        if any(sorted(t.schedule_times or []) == v["schedule_times"] for t in same):
            self.report.skipped(sheet.name)
            return
        if v["end_date"] and v["end_date"] < timezone.localdate():
            self.report.warning(sheet.name, row, "terapia già terminata: viene importata come terminata", "Fine")
        therapy = Therapy(
            guest=guest, drug=drug, dose=v["dose"], dose_unit=v["dose_unit"],
            stock_units_per_dose=v["stock_units_per_dose"], route=v["route"], schedule_times=v["schedule_times"],
            every_n_days=v["every_n_days"] or 1, start_date=v["start_date"], end_date=v["end_date"],
            instructions=v["instructions"], is_sos=v["is_sos"], sos_indication=v["sos_indication"],
            sos_max_daily=v["sos_max_daily"], sos_min_interval_hours=v["sos_min_interval_hours"],
            prescriber_name=v["prescriber_name"], prescription_date=v["prescription_date"] or v["start_date"],
        )
        if therapy.end_date and therapy.end_date < timezone.localdate():
            therapy.status = Therapy.Status.ENDED
        labels = {c.key: c.header for c in sheet.columns}
        if not _clean_or_report(therapy, self.report, sheet.name, row, exclude=["prescription_document"],
                                labels=labels):
            return
        warnings = check_therapy_conflicts(guest, drug)
        for warning in warnings:
            self.report.warning(sheet.name, row, f"{guest}: {warning.message}")
        if warnings:
            therapy.conflict_acknowledged = "Importazione dal cartaceo: " + " | ".join(w.message for w in warnings)
        therapy.save()
        self.report.created(sheet.name)

    def import_magazzino(self, sheet, row, v):
        guest = None
        if v["owner_cf"]:
            guest = self.guest_for(sheet, row, v["owner_cf"], "Codice fiscale ospite (farmaco personale)")
            if guest is None:
                return
        drug = self.drug_for(sheet, row, v)
        if drug is None:
            return
        if v["quantity"] <= 0:
            self.report.error(sheet.name, row, "la quantità deve essere maggiore di zero", "Quantità")
            return
        if StockLot.objects.filter(drug=drug, guest=guest, lot_number__iexact=v["lot_number"],
                                   expiry_date=v["expiry_date"]).exists():
            self.report.skipped(sheet.name)
            return
        if v["expiry_date"] < timezone.localdate():
            self.report.warning(sheet.name, row, "lotto già scaduto: verrà segnalato tra gli avvisi di magazzino",
                                "Scadenza")
        load_lot(drug, v["lot_number"], v["expiry_date"], v["quantity"], self.user, guest=guest,
                 supplier=v["supplier"], notes="Giacenza iniziale (importazione dal cartaceo)")
        self.report.created(sheet.name)


def run_workbook_import(data, user, commit=False, filename=""):
    report = ImportReport()
    parsed = read_workbook(data, report)
    if not parsed:
        return report
    with audit_context(user=user, source="import", reason=f"Importazione FASAS cartaceo ({filename})".strip()):
        try:
            with transaction.atomic():
                _Importer(parsed, user, report).run()
                if not commit or report.errors:
                    raise _Rollback
        except _Rollback:
            pass
        else:
            report.committed = True
    return report


# --- File in attesa di conferma -------------------------------------------------------------

def store_workbook(data, filename, user, report):
    return ImportBatch.objects.create(
        kind=ImportBatch.Kind.WORKBOOK, filename=filename[:255], sha256=sha256(data), created_by=user,
        payload=[store_encrypted(data, "imports")], summary=report.as_dict(),
    )


def workbook_bytes(batch):
    return read_encrypted(batch.payload[0]) if batch.payload else b""


def close_workbook(batch, status, report=None):
    """Il file caricato si cancella appena l'importazione è conclusa o annullata."""
    for name in batch.payload:
        default_storage.delete(name)
    batch.payload = []
    batch.status = status
    batch.completed_at = timezone.now()
    if report is not None:
        batch.summary = report.as_dict()
    batch.save()


def parse_date_arg(text):
    """Data da riga di comando (GG/MM/AAAA o AAAA-MM-GG)."""
    from .spec import Column, parse_date

    return parse_date(text, Column("data", "data")) if text else None
