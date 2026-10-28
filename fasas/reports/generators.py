"""Dati dei report periodici: ogni generatore restituisce titolo, colonne, righe e larghezze relative."""
import datetime
from dataclasses import dataclass, field
from io import BytesIO

from django.utils import timezone

from .models import ReportFormat, ReportType


@dataclass
class ReportData:
    title: str
    headers: list
    rows: list = field(default_factory=list)
    widths: list | None = None


def _period(qs, start, end, field_name="shift_date"):
    return qs.filter(**{f"{field_name}__gte": start, f"{field_name}__lte": end})


def _guest_filter(qs, guest, field_name="guest"):
    return qs.filter(**{field_name: guest}) if guest else qs


def _t(dt):
    return timezone.localtime(dt).strftime("%d/%m/%Y %H:%M") if dt else ""


def _who(user):
    return str(user) if user else ""


def daily_register(start, end, guest=None):
    from fasas.daily.models import DailyRecord, DiaryEntry, EvacuationRecord

    rows = []
    for r in _guest_filter(_period(DailyRecord.objects.select_related("guest", "created_by"), start, end), guest):
        detail = " – ".join(x for x in [r.get_meal_display(), r.get_intake_display(), r.description] if x)
        rows.append([_t(r.occurred_at), r.get_shift_display(), str(r.guest), r.get_category_display(), detail,
                     _who(r.created_by)])
    for e in _guest_filter(_period(EvacuationRecord.objects.select_related("guest", "created_by"), start, end), guest):
        detail = "SI" if e.outcome else "NO"
        if e.consistency:
            detail += f" – {e.get_consistency_display()}"
        rows.append([_t(e.occurred_at), e.get_shift_display(), str(e.guest), "Evacuazione", f"{detail} {e.notes}".strip(),
                     _who(e.created_by)])
    for d in _guest_filter(_period(DiaryEntry.objects.select_related("guest", "created_by"), start, end), guest):
        rows.append([_t(d.occurred_at), d.get_shift_display(), str(d.guest),
                     "Diario" + (" (rilevante)" if d.is_relevant else ""), d.text, _who(d.created_by)])
    rows.sort(key=lambda r: (datetime.datetime.strptime(r[0], "%d/%m/%Y %H:%M"), r[2]))
    return ReportData("Registro giornaliero", ["Data/ora", "Turno", "Ospite", "Categoria", "Dettaglio", "Operatore"],
                      rows, [1.3, 0.8, 1.5, 1.2, 4.5, 1.3])


def medications(start, end, guest=None):
    from fasas.clinical.models import TherapyAdministration

    qs = _guest_filter(_period(TherapyAdministration.objects.select_related("guest", "therapy__drug", "created_by",
                                                                           "validated_by"), start, end), guest)
    rows = []
    for a in qs.order_by("occurred_at"):
        scheduled = f"{a.scheduled_date:%d/%m} {a.scheduled_time:%H:%M}" if a.scheduled_time else "SOS"
        validation = a.get_validation_status_display() if a.is_sos else ""
        if a.validated_by:
            validation += f" ({a.validated_by})"
        rows.append([_t(a.occurred_at), str(a.guest), str(a.therapy.drug), a.dose_given, scheduled,
                     a.get_status_display(), a.sos_reason or a.notes, validation, _who(a.created_by)])
    return ReportData("Farmaci somministrati", ["Data/ora", "Ospite", "Farmaco", "Dose", "Programmata", "Esito",
                                                "Motivo/Note", "Validazione SOS", "Operatore"],
                      rows, [1.2, 1.4, 1.5, 0.8, 0.9, 1.0, 2.2, 1.3, 1.2])


def events(start, end, guest=None):
    from fasas.clinical.models import ClinicalEvent, Restraint

    rows = []
    for e in _guest_filter(_period(ClinicalEvent.objects.select_related("guest", "created_by"), start, end), guest):
        rows.append([_t(e.occurred_at), str(e.guest), e.get_event_type_display(), e.get_severity_display(),
                     e.description, e.actions_taken, _who(e.created_by)])
    for r in _guest_filter(_period(Restraint.objects.select_related("guest", "responsible_operator"), start, end), guest):
        rows.append([_t(r.occurred_at), str(r.guest), f"Contenzione: {r.get_restraint_type_display()}",
                     r.duration_display, r.reason, f"Autorizzata da {r.authorizing_doctor_display}",
                     _who(r.responsible_operator)])
    rows.sort(key=lambda r: datetime.datetime.strptime(r[0], "%d/%m/%Y %H:%M"))
    return ReportData("Eventi sanitari e contenzioni", ["Data/ora", "Ospite", "Tipo", "Gravità/durata", "Descrizione",
                                                        "Interventi", "Operatore"], rows, [1.2, 1.4, 1.5, 1, 3, 2.2, 1.2])


def parameters(start, end, guest=None):
    from fasas.clinical.models import ParameterReading

    qs = _guest_filter(_period(ParameterReading.objects.select_related("guest", "parameter_type", "created_by"),
                               start, end), guest)
    rows = [[_t(r.occurred_at), str(r.guest), r.parameter_type.name, r.display_value,
             "FUORI RANGE" if r.out_of_range else "", r.notes, _who(r.created_by)] for r in qs.order_by("occurred_at")]
    return ReportData("Parametri sanitari", ["Data/ora", "Ospite", "Parametro", "Valore", "Allerta", "Note", "Operatore"],
                      rows, [1.2, 1.5, 1.2, 1, 0.9, 2.5, 1.2])


def protocols(start, end, guest=None):
    from fasas.protocols.models import ProtocolActivation

    qs = ProtocolActivation.objects.select_related("guest", "protocol", "resolved_by").filter(
        triggered_at__date__gte=start, triggered_at__date__lte=end
    )
    qs = _guest_filter(qs, guest)
    rows = []
    for a in qs.order_by("triggered_at").prefetch_related("action_logs__done_by"):
        actions = "; ".join(f"{log.action_label} ({_who(log.done_by)} {_t(log.done_at)})" for log in a.action_logs.all())
        rows.append([_t(a.triggered_at), str(a.guest), a.protocol.name, a.trigger_summary, actions,
                     a.get_status_display(), a.resolution_notes, _who(a.resolved_by)])
    return ReportData("Protocolli attivati", ["Attivato", "Ospite", "Protocollo", "Motivo", "Azioni eseguite", "Stato",
                                              "Esito", "Chiuso da"], rows, [1.1, 1.3, 1.3, 2, 2.6, 0.9, 1.6, 1.1])


GENERATORS = {
    ReportType.DAILY_REGISTER: daily_register,
    ReportType.MEDICATIONS: medications,
    ReportType.EVENTS: events,
    ReportType.PARAMETERS: parameters,
    ReportType.PROTOCOLS: protocols,
}


def render_pdf(data, start, end, guest=None, user=None):
    from .pdf import Document

    subtitle = f"Periodo {start:%d/%m/%Y} – {end:%d/%m/%Y}"
    if guest:
        subtitle += f" · Ospite: {guest}"
    doc = Document(data.title, subtitle, orientation="landscape", generated_by=user)
    doc.text(f"{len(data.rows)} registrazioni.", "small")
    doc.table(data.headers, data.rows, data.widths)
    return doc.build()


def _safe_cell(value):
    """Impedisce che testo inserito dagli utenti venga interpretato come formula da Excel."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def render_xlsx(data, start, end, guest=None, user=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = data.title[:31]
    ws.append([data.title])
    ws["A1"].font = Font(bold=True, size=14)
    info = f"Periodo {start:%d/%m/%Y} – {end:%d/%m/%Y}"
    if guest:
        info += f" · Ospite: {guest}"
    ws.append([info + f" · Generato il {timezone.localtime():%d/%m/%Y %H:%M}" + (f" da {user}" if user else "")])
    ws.append([])
    ws.append(data.headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
    for row in data.rows:
        ws.append([_safe_cell(v) for v in row])
    widths = data.widths or [1] * len(data.headers)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(12, int(width * 14))
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(data.headers))}{ws.max_row}"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render(report_type, fmt, start, end, guest=None, user=None):
    data = GENERATORS[report_type](start, end, guest)
    if fmt == ReportFormat.XLSX:
        return data, render_xlsx(data, start, end, guest, user)
    return data, render_pdf(data, start, end, guest, user)
