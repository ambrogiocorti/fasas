"""Esportazioni per la rendicontazione: presenze mensili ed elenco prospetti (Excel), prospetto singolo (PDF)."""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fasas.reports.pdf import Document, p, table

from .models import AttendanceStatus, BillingSettings
from .services import attendance, month_days

MONTHS = ["", "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre",
          "ottobre", "novembre", "dicembre"]
FILLS = {"P": "DCEFE3", "H": "F9D7D3", "F": "FCEBC7", "A": "E4E7EB", "-": "FFFFFF"}


def month_label(first):
    return f"{MONTHS[first.month]} {first.year}"


def _header(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FAD")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def attendance_workbook(first):
    month, rows = attendance(first)
    days = month_days(first)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Presenze {first:%m-%Y}"
    ws.append([f"Presenze {month_label(first)} – {'mese chiuso' if month and month.is_closed else 'provvisorie (mese aperto)'}"])
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, ["Cognome e nome", "Codice fiscale"] + [str(d.day) for d in days]
            + ["Presenze", "Ricovero", "Famiglia", "Altre assenze", "Giorni in carico"])
    for guest, values, c, _ in rows:
        ws.append([f"{guest.last_name} {guest.first_name}", guest.fiscal_code] + [values.get(d.day, "-") for d in days]
                  + [c["present"], c["hospital"], c["family"], c["absent"], sum(c.values())])
        for col in range(3, 3 + len(days)):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.fill = PatternFill("solid", fgColor=FILLS.get(cell.value, "FFFFFF"))
            cell.alignment = Alignment(horizontal="center")
    totals = ["Totale", ""] + [""] * len(days) + [
        sum(r[2]["present"] for r in rows), sum(r[2]["hospital"] for r in rows), sum(r[2]["family"] for r in rows),
        sum(r[2]["absent"] for r in rows), sum(sum(r[2].values()) for r in rows)]
    ws.append(totals)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 19
    for col in range(3, 3 + len(days)):
        ws.column_dimensions[get_column_letter(col)].width = 3.6
    for col in range(3 + len(days), 8 + len(days)):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.freeze_panes = "C3"
    legend = wb.create_sheet("Legenda")
    legend.append(["Codice", "Significato"])
    for code, label in AttendanceStatus.choices:
        legend.append([code, label])
    rules = BillingSettings.get()
    legend.append([])
    legend.append(["Regole", f"Giorno di ingresso {'conteggiato' if rules.count_admission_day else 'non conteggiato'}; "
                             f"giorno di dimissione {'conteggiato' if rules.count_discharge_day else 'non conteggiato'}; "
                             f"giorno di uscita {'presente' if rules.departure_day_present else 'assente'}; "
                             f"giorno di rientro {'presente' if rules.return_day_present else 'assente'}."])
    if month and month.notes:
        legend.append(["Note", month.notes])
    legend.column_dimensions["B"].width = 90
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def statements_workbook(month):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Prospetti {month.month:%m-%Y}"
    _header(ws, ["Numero", "Emesso il", "Scadenza", "Stato", "Ospite", "Codice fiscale ospite", "Pagante",
                 "Intestatario", "CF / P.IVA intestatario", "Totale €", "Incassato €", "Residuo €"])
    for s in month.statements.select_related("guest").prefetch_related("payments").order_by("number", "guest__last_name"):
        ws.append([s.number or "bozza", s.issue_date, s.due_date, s.get_status_display(),
                   f"{s.guest.last_name} {s.guest.first_name}", s.guest.fiscal_code, s.get_payer_display(),
                   s.payer_name, s.payer_tax_code, float(s.total), float(s.paid_amount), float(s.balance)])
        for col in (2, 3):
            ws.cell(row=ws.max_row, column=col).number_format = "DD/MM/YYYY"
        for col in (10, 11, 12):
            ws.cell(row=ws.max_row, column=col).number_format = "#,##0.00"
    widths = [12, 12, 12, 10, 24, 19, 26, 30, 18, 12, 12, 12]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    detail = wb.create_sheet("Righe")
    _header(detail, ["Numero", "Ospite", "Intestatario", "Voce", "Presenza", "Giorni", "Tariffa €", "%", "Importo €"])
    for s in month.statements.select_related("guest").exclude(status="cancelled"):
        for line in s.lines:
            detail.append([s.number or "bozza", f"{s.guest.last_name} {s.guest.first_name}", s.payer_name,
                           line["description"], line["status_label"], line["days"], float(Decimal(line["rate"])),
                           line["percent"], float(Decimal(line["amount"]))])
    for index, width in enumerate([12, 24, 30, 22, 34, 8, 10, 6, 12], start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def euro(value):
    text = f"{Decimal(value):,.2f}"
    return "€ " + text.replace(",", "X").replace(".", ",").replace("X", ".")


def statement_pdf(statement, user):
    rules = BillingSettings.get()
    title = f"Prospetto retta {statement.number or '(bozza)'} – {month_label(statement.month.month)}"
    doc = Document(title, generated_by=user, subtitle=rules.payee_name or "")
    payee = [x for x in (rules.payee_name, rules.payee_address,
                         f"CF/P.IVA {rules.payee_tax_code}" if rules.payee_tax_code else "") if x]
    doc.key_values([
        ("Ente gestore", " · ".join(payee) or doc.community.name),
        ("Intestatario", " · ".join(x for x in (statement.payer_name, statement.payer_address,
                                                 f"CF/P.IVA {statement.payer_tax_code}" if statement.payer_tax_code else "")
                                     if x)),
        ("Ospite", f"{statement.guest.last_name} {statement.guest.first_name} – {statement.guest.fiscal_code}"),
        ("Periodo", month_label(statement.month.month)),
        ("Emesso il", statement.issue_date), ("Scadenza", statement.due_date),
    ])
    rows = [(line["description"], line["status_label"], line["days"], euro(line["rate"]), f"{line['percent']}%",
             euro(line["amount"])) for line in statement.lines]
    doc.table(["Voce", "Presenza", "Giorni", "Tariffa giornaliera", "%", "Importo"], rows, [3, 4, 1, 2, 1, 2])
    doc.add(p(f"Totale: {euro(statement.total)}", "h1"))
    if statement.paid_amount:
        doc.text(f"Già incassato: {euro(statement.paid_amount)} · residuo: {euro(statement.balance)}")
    if rules.iban:
        doc.text(f"Pagamento con bonifico su IBAN {rules.iban} indicando il numero del prospetto.")
    if rules.statement_notes:
        doc.text(rules.statement_notes, "small")
    if statement.status == "cancelled":
        doc.add(p(f"PROSPETTO ANNULLATO: {statement.cancel_reason}", "h2"))
    doc.add(table(["Legenda presenze"], [[label] for code, label in AttendanceStatus.choices if code != "-"],
                  available_width=doc.width))
    return doc.build()
