import io
import json

from django.contrib import messages
from django.contrib.auth.decorators import permission_required
from django.http import HttpResponse, HttpResponseBadRequest
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fasas.billing.exports import month_label
from fasas.billing.services import month_days, month_start, parse_month, shift_month

from . import services
from .forms import (
    PAINT_CODES,
    DecisionForm,
    LeaveForm,
    RotaSettingsForm,
    ShiftChangeForm,
    StaffProfileFormSet,
)
from .models import Code, LeaveRequest, RequestStatus, RotaEntry, RotaMonth, ShiftChange

CODE_FILLS = {"M": "DCEBFA", "P": "FCEBC7", "N": "D9D4F4", "R": "FFFFFF", "F": "DCEFE3", "PE": "E4F2EA",
              "MA": "F9D7D3", "FO": "E4E7EB"}


def _first(request):
    return parse_month(request.GET.get("mese") or request.POST.get("mese"))


def _url(name, first):
    return f"{reverse(name)}?mese={first:%Y-%m}"


@permission_required("staffing.view_rotamonth", raise_exception=True)
def rota(request):
    first = _first(request)
    grid = services.month_grid(first)
    return render(request, "staffing/rota.html", {
        **grid, "first": first, "label": month_label(first), "prev": shift_month(first, -1),
        "next": shift_month(first, 1), "codes": PAINT_CODES, "today": timezone.localdate(),
        "can_edit": request.user.has_perm("staffing.change_rotaentry"),
        "weekend": {d for d in grid["days"] if d.weekday() >= 5},
    })


@require_POST
@permission_required("staffing.change_rotaentry", raise_exception=True)
def save_changes(request):
    first = _first(request)
    try:
        changes = json.loads(request.POST.get("changes") or "[]")
        if not isinstance(changes, list):
            raise ValueError
    except ValueError:
        return HttpResponseBadRequest("Modifiche non valide")
    try:
        count = services.apply_changes(changes, request.user)
    except services.RotaError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"{count} modifiche salvate.")
    return redirect(_url("staffing:rota", first))


@require_POST
@permission_required("staffing.change_rotaentry", raise_exception=True)
def repeat(request):
    first = _first(request)
    created = services.repeat_cycle(first, int(request.POST.get("weeks") or 4))
    messages.success(request, f"{created} turni copiati dallo schema delle settimane precedenti: controllare e correggere.")
    return redirect(_url("staffing:rota", first))


@require_POST
@permission_required("staffing.publish_rota", raise_exception=True)
def publish_view(request):
    first = _first(request)
    services.publish(first, request.user, request.POST.get("notes", "").strip())
    messages.success(request, f"Turni di {month_label(first)} pubblicati: il personale riceve una notifica.")
    return redirect(_url("staffing:rota", first))


@permission_required("staffing.view_rotamonth", raise_exception=True)
def rota_excel(request):
    first = _first(request)
    grid = services.month_grid(first)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Turni {first:%m-%Y}"
    ws.append([f"Turni del personale – {month_label(first)}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Operatore"] + [d.day for d in grid["days"]] + ["Ore", "Ore attese", "Ferie", "Notti"])
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in grid["rows"]:
        ws.append([str(row["user"])] + [c["code"] for c in row["cells"]]
                  + [float(row["hours"]), float(row["expected"]), row["holidays"], row["nights"]])
        for col, cell in enumerate(row["cells"], start=2):
            if cell["code"]:
                ws.cell(row=ws.max_row, column=col).fill = PatternFill("solid", fgColor=CODE_FILLS.get(cell["code"], "FFFFFF"))
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(horizontal="center")
    for code, cells in grid["coverage"].items():
        ws.append([f"Copertura {Code(code).label.lower()}"] + [c["count"] for c in cells])
    ws.column_dimensions["A"].width = 26
    for col in range(2, 2 + len(grid["days"])):
        ws.column_dimensions[get_column_letter(col)].width = 4
    legend = wb.create_sheet("Legenda")
    for code, text in Code.choices:
        legend.append([code, text])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = content_disposition_header(True, f"turni-{first:%Y-%m}.xlsx")
    return response


@permission_required("staffing.view_rotamonth", raise_exception=True)
def my_rota(request):
    user = request.user
    today = timezone.localdate()
    months = []
    for first in (month_start(today), shift_month(month_start(today), 1)):
        month = RotaMonth.objects.filter(month=first).first()
        entries = {e.date: e for e in RotaEntry.objects.filter(user=user, date__gte=first, date__lte=month_days(first)[-1])}
        months.append({"first": first, "label": month_label(first), "month": month, "offset": range(first.weekday()),
                       "days": [(d, entries.get(d)) for d in month_days(first)] if month and month.is_published else []})
    can_request = user.has_perm("staffing.add_leaverequest")
    leave_form = LeaveForm(request.POST or None, prefix="leave") if can_request else None
    colleagues = [u for u, _ in services.staff()]
    change_form = ShiftChangeForm(request.POST or None, prefix="change", user=user, colleagues=colleagues) \
        if user.has_perm("staffing.add_shiftchange") else None
    if request.method == "POST":
        try:
            if "leave-kind" in request.POST and leave_form is not None and leave_form.is_valid():
                d = leave_form.cleaned_data
                services.request_leave(user, d["kind"], d["start"], d["end"], d["note"])
                messages.success(request, "Richiesta inviata al coordinatore.")
                return redirect("staffing:my")
            if "change-kind" in request.POST and change_form is not None and change_form.is_valid():
                d = change_form.cleaned_data
                services.propose_change(user, d["kind"], d["date"], d["colleague"], d["note"])
                messages.success(request, f"Proposta inviata a {d['colleague']}.")
                return redirect("staffing:my")
        except services.RotaError as exc:
            messages.error(request, str(exc))
    return render(request, "staffing/my.html", {
        "months": months, "today": today, "leave_form": leave_form, "change_form": change_form,
        "leaves": LeaveRequest.objects.filter(user=user)[:20],
        "incoming": ShiftChange.objects.filter(colleague=user, status=RequestStatus.PENDING).select_related("requester"),
        "outgoing": ShiftChange.objects.filter(requester=user).select_related("colleague")[:20],
        "upcoming": services.upcoming(user, 14),
    })


@require_POST
@permission_required("staffing.add_leaverequest", raise_exception=True)
def cancel_leave(request, pk):
    leave = get_object_or_404(LeaveRequest, pk=pk, user=request.user, status=RequestStatus.PENDING)
    leave.status = RequestStatus.CANCELLED
    leave.save()
    messages.info(request, "Richiesta ritirata.")
    return redirect("staffing:my")


@require_POST
@permission_required("staffing.add_shiftchange", raise_exception=True)
def respond_change_view(request, pk):
    change = get_object_or_404(ShiftChange, pk=pk)
    try:
        services.respond_change(change, request.user, request.POST.get("accept") == "1")
    except services.RotaError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Risposta inviata.")
    return redirect("staffing:my")


@permission_required("staffing.change_leaverequest", raise_exception=True)
def requests_view(request):
    pending = LeaveRequest.objects.filter(status=RequestStatus.PENDING).select_related("user").order_by("start")
    leaves = [{"leave": leave, "impact": services.coverage_impact(leave.user, services._days(leave.start, leave.end),
                                                                   leave.kind)} for leave in pending]
    changes = ShiftChange.objects.filter(status__in=[RequestStatus.PENDING, RequestStatus.ACCEPTED]) \
        .select_related("requester", "colleague").order_by("date")
    history = list(LeaveRequest.objects.exclude(status=RequestStatus.PENDING).select_related("user", "decided_by")[:15])
    return render(request, "staffing/requests.html", {
        "leaves": leaves, "changes": changes, "history": history, "form": DecisionForm(),
        "changes_history": ShiftChange.objects.exclude(status__in=[RequestStatus.PENDING, RequestStatus.ACCEPTED])
        .select_related("requester", "colleague")[:15],
    })


@require_POST
@permission_required("staffing.change_leaverequest", raise_exception=True)
def decide_leave_view(request, pk):
    leave = get_object_or_404(LeaveRequest, pk=pk)
    try:
        services.decide_leave(leave, request.user, request.POST.get("approve") == "1", request.POST.get("note", ""))
    except services.RotaError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"Richiesta di {leave.user} {'approvata' if leave.status == 'approved' else 'respinta'}.")
    return redirect("staffing:requests")


@require_POST
@permission_required("staffing.change_shiftchange", raise_exception=True)
def decide_change_view(request, pk):
    change = get_object_or_404(ShiftChange, pk=pk)
    try:
        services.decide_change(change, request.user, request.POST.get("approve") == "1", request.POST.get("note", ""))
    except services.RotaError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Decisione registrata.")
    return redirect("staffing:requests")


@permission_required("staffing.change_rotasettings", raise_exception=True)
def settings_view(request):
    services.staff()  # crea i profili mancanti
    rules = services.RotaSettings.get()
    form = RotaSettingsForm(request.POST or None, instance=rules, prefix="rules")
    formset = StaffProfileFormSet(request.POST or None, prefix="staff",
                                  queryset=services.StaffProfile.objects.select_related("user")
                                  .filter(user__is_active=True).order_by("order", "user__last_name"))
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        form.save()
        formset.save()
        messages.success(request, "Regole e profili salvati.")
        return redirect("staffing:settings")
    return render(request, "staffing/settings.html", {"form": form, "formset": formset})


def pending_requests_count():
    return LeaveRequest.objects.filter(status=RequestStatus.PENDING).count() + \
        ShiftChange.objects.filter(status=RequestStatus.ACCEPTED).count()

